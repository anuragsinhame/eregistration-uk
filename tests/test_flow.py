"""Mock-based flow test for uk_ereg_gui.Worker - no real browser, webview or network needed.

Run:  python tests/test_flow.py

A FakePage stands in for Playwright: login with captcha, the menu, Buyer/Seller pages, cascading
dropdowns, the custom report (captcha per search, two result pages, cancel) and the Excel export.
"""
import contextlib
import json
import os
import queue
import re
import shutil
import sys
import tempfile
import time
import types
from unittest.mock import patch
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent   # the project folder (uk_ereg_gui.py lives there)

# ---------------------------------------------------------------- stub third-party modules
sys.modules["anthropic"] = types.ModuleType("anthropic")
webview = types.ModuleType("webview")
class FileDialog: FOLDER = "FOLDER"; SAVE = "SAVE"
webview.FileDialog = FileDialog
sys.modules["webview"] = webview
pw_pkg = types.ModuleType("playwright"); sys.modules["playwright"] = pw_pkg
pw_sync = types.ModuleType("playwright.sync_api"); sys.modules["playwright.sync_api"] = pw_sync
class PWError(Exception): pass
class PWTimeout(PWError): pass
pw_sync.Error = PWError; pw_sync.TimeoutError = PWTimeout; pw_sync.Page = object

SEARCH_URL = "https://online.eregistrationukgov.in/E_Search/Default2.aspx"
BUYER_URL = "https://online.eregistrationukgov.in/E_Search/BuyerWise.aspx"
SELLER_URL = "https://online.eregistrationukgov.in/E_Search/SellerWise.aspx"
CAPTCHA_ANSWER = "AB12"
PNG = b"\x89PNG\r\n\x1a\n fake png bytes"


class FakeLocator:
    def __init__(self, page, sel):
        self.page, self.sel = page, sel
    @property
    def first(self): return self
    def nth(self, i): return FakeLocator(self.page, f"{self.sel}[{i}]")
    def filter(self, **k): return FakeLocator(self.page, f"{self.sel}[has_text={k.get('has_text')}]")
    def count(self): return 0 if "missing" in self.sel else (2 if "download" in self.sel else 1)
    def wait_for(self, **k): pass
    def evaluate(self, js, *a):
        if "naturalWidth" in js: return True
        if "el.value" in js: return {"#MainContent_btnSearch": "Search", "#MainContent_btnReset": "Reset"}.get(self.sel, self.sel)
        if "s.options" in js: return [["1", "ALMORA"], ["2", "BAGESHWAR"]]
        return None
    def screenshot(self, **k): return PNG
    def fill(self, v): self.page.filled[self.sel] = v
    def input_value(self): return self.page.filled.get(self.sel, "")
    def select_option(self, **k):
        self.page.filled[self.sel] = k; self.page.postbacks += 1
        if "Year" in self.sel: self.page.year = k.get("label") or k.get("value")
        if "District" in self.sel and k.get("label") and k["label"] not in ("ALMORA", "BAGESHWAR"): raise PWError("no such option")
    def set_checked(self, v): self.page.filled[self.sel] = v
    def check(self): self.page.filled[self.sel] = True
    def hover(self, **k): pass
    def click(self, **k):
        self.page.clicks.append(self.sel)
        s = str(self.sel)
        if "download" in s:
            self.page.pending_download = f"doc_{len(self.page.clicks)}.pdf"
        elif s == "#MainContent_btnSearch":
            if self.page.filled.get("#MainContent_txtCaptcha2", "") == CAPTCHA_ANSWER:
                self.page.state = "results"; self.page.body = "Results\n"; self.page.result_page = 1
                self.page.searches.append((self.page.role, self.page.year, self.page.filled.get("#MainContent_txtBuyer")))
                if self.page.role == "khasra":
                    self.page.khasra_searches.append((self.page.year, self.page.filled.get("#MainContent_txtKhasra")))
            else:
                self.page.body = "Invalid Captcha, please try again\n"
        elif "Buyer Wise" in s or "Buyer" in s:
            self.page.url = BUYER_URL; self.page.state = "search"; self.page.role = "buyer"
        elif "Khasra" in s:
            self.page.url = "https://online.eregistrationukgov.in/E_Search/frm_index_khasra.aspx"
            self.page.state = "search"; self.page.role = "khasra"
        elif "Seller" in s:
            self.page.url = SELLER_URL; self.page.state = "search"; self.page.role = "seller"
        elif "data-ereg-next" in s:
            self.page.result_page += 1
        elif "data-ereg-dl" in s:
            self.page.pending_download = f"{self.page.marked}.pdf"
            self.page.doc_clicks.append((self.page.role, self.page.year, self.page.result_page, self.page.marked))


class FakeDownload:
    def __init__(self, name): self.suggested_filename = name
    def save_as(self, target): Path(target).write_bytes(b"%PDF-1.4 fake")


class FakePage:
    def __init__(self):
        self.url = "about:blank"; self.filled = {}; self.clicks = []; self.postbacks = 0
        self.state = "login"; self.body = ""; self.pending_download = None; self.login_attempts = 0
        self.role = "buyer"; self.year = "2023"; self.result_page = 1; self.searches = []
        self.khasra_searches = []
        self.marked = None; self.doc_clicks = []
        self.context = types.SimpleNamespace(pages=[self])
    def set_default_timeout(self, ms): pass
    def goto(self, url, **k):
        self.url = url
        if "Default2" in url: self.state = "home"
    def fill(self, sel, v): self.filled[sel] = v
    def click(self, sel, **k):
        self.clicks.append(sel)
        if sel == "#MainContent_btnLogin":
            self.login_attempts += 1
            if self.filled.get("#MainContent_txtCaptcha") == CAPTCHA_ANSWER:
                self.url = SEARCH_URL; self.state = "home"; self.body = "Welcome\n"
            else:
                self.body = "Invalid Captcha\n"
    def wait_for_load_state(self, *a, **k): pass
    def wait_for_timeout(self, ms): pass
    def on(self, event, handler): pass
    main_frame = None
    def inner_text(self, sel): return self.body
    def locator(self, sel): return FakeLocator(self, sel)
    def get_by_text(self, text, **k): return FakeLocator(self, f"text={text}")
    def get_by_role(self, *a, **k): return FakeLocator(self, "role")
    def screenshot(self, **k): return b"\xff\xd8\xff fakejpeg"
    def evaluate(self, js, *a):
        if "scrollHeight" in js: return 1000
        if "data-ereg-next" in js: return {"status": "ok", "current": 1, "next": "2"} if self.result_page == 1 else {"status": "last", "current": 2}
        if "data-ereg-dl" in js:
            cells = a[0][1]
            rows = [r for r in self.scan()["tables"][0]["rows"] if len(r) == 3]
            if cells not in rows: return {"status": "no-row"}
            self.marked = cells[1]; return {"status": "ok", "row": cells}
        if "entries|records" in js: return None
        return self.scan()
    @contextlib.contextmanager
    def expect_download(self, **k):
        holder = types.SimpleNamespace(value=None)
        yield holder
        if not self.pending_download: raise PWTimeout("no download")
        holder.value = FakeDownload(self.pending_download); self.pending_download = None
    def form_controls(self):
        who = "Buyer" if self.role == "buyer" else "Seller"
        controls = [
            {"kind": "select", "selector": "#MainContent_ddlDistrict", "label": "District", "value": "1",
             "options": [{"value": "0", "label": "--Select--", "selected": False}, {"value": "1", "label": "ALMORA", "selected": True}, {"value": "2", "label": "BAGESHWAR", "selected": False}]},
            {"kind": "select", "selector": "#MainContent_ddlSRO", "label": "Sub-Registrar Office", "value": "21",
             "options": [{"value": "21", "label": "BAGESHWAR", "selected": True}, {"value": "22", "label": "KAPKOT", "selected": False}]},
            {"kind": "select", "selector": "#MainContent_ddlYear", "label": "Registration Year", "value": "2023",
             "options": [{"value": y, "label": y, "selected": y == self.year} for y in ("2021", "2022", "2023", "2024")]},
            {"kind": "textarea", "selector": "#MainContent_txtBuyer", "label": f"Enter {who} Name", "value": ""},
            {"kind": "text", "selector": "#tblSearch_filter", "label": "Search...", "value": ""},
            {"kind": "button", "selector": "#MainContent_btnSearch", "label": "Search"},
            {"kind": "button", "selector": "#MainContent_btnReset", "label": "Reset"}]
        if self.role == "khasra":
            controls[3] = {"kind": "text", "selector": "#MainContent_txtKhasra", "label": "Khasra No.", "value": ""}
        return controls
    def scan(self):
        base = {"url": self.url, "title": "E-Search", "heading": "", "messages": [], "links": [
            {"text": "Buyer Wise", "hover": "Search By Party Name", "href": "#", "visible": False}], "tables": [], "controls": [], "captcha": None}
        if self.state in ("search", "results"):
            base["controls"] = self.form_controls()
            base["captcha"] = {"img": "#MainContent_imgCaptcha2", "input": "#MainContent_txtCaptcha2"}
        if self.state == "results":
            tag = f"{self.role}-{self.year}-p{self.result_page}"
            rows = [["1", f"{tag}-A", "Download Document"], ["2", f"{tag}-B", "Download Document"]] if self.result_page == 1 else [["3", f"{tag}-C", "Download Document"]]
            rows.append(["1 2"])   # pager row (colspan) must be dropped
            base["tables"] = [{"id": "gv", "index": 3, "header": ["Sr", "Buyer", "Doc"], "rows": rows, "total_rows": len(rows), "downloads": len(rows) - 1}]
        return base


class FakeBrowser:
    def __init__(self): self.page = FakePage(); self.closed = False
    def new_context(self, **k): return types.SimpleNamespace(new_page=lambda: self.page)
    def close(self): self.closed = True

class FakePW:
    def __init__(self):
        self.browser = FakeBrowser(); self.stopped = False; self.fail_launch_once = False; self.launches = 0
        self.chromium = types.SimpleNamespace(launch=self._launch)
    def _launch(self, **k):
        self.launches += 1
        if self.fail_launch_once:
            self.fail_launch_once = False
            raise PWError("BrowserType.launch: Executable doesn't exist at /x/chrome\nPlease run the following command to download new browsers:\n    playwright install")
        return self.browser
    def stop(self): self.stopped = True

FAKE_PW = FakePW()
pw_sync.sync_playwright = lambda: types.SimpleNamespace(start=lambda: FAKE_PW)

# ---------------------------------------------------------------- import the module under test
sys.path.insert(0, str(ROOT))
import uk_ereg_gui as gui  # noqa: E402
_SETTINGS_TMP = Path(tempfile.mkdtemp())
gui.SETTINGS_FILE = _SETTINGS_TMP / "settings.json"

events = queue.Queue()
class RecordingUI:
    def emit(self, ev):
        json.dumps(ev)                                   # must be JSON-serialisable
        events.put(ev)
        kind = ev.get("type")
        out = sys.__stdout__
        if kind == "page": out.write(f"  <- page[{ev.get('mode')}]: {len(ev['controls'])} controls, captcha={bool(ev.get('captcha'))}, tables={len(ev['tables'])}\n")
        elif kind in ("log", "error"): out.write(f"  <- {kind}: {ev['message']}\n")
        elif kind == "captcha": out.write(f"  <- captcha prompt: {ev['prompt']!r} allow_skip={ev['allow_skip']}\n")
        elif kind == "report_progress": out.write(f"  <- progress: {ev['role']} {ev['year']} {ev['message']} (total {ev['total']})\n")
        elif kind == "report_done": out.write(f"  <- report_done: partial={ev['partial']} " + ", ".join(f"{r}={len(s['rows'])}" for r, s in ev['sheets'].items()) + "\n")
        elif kind in ("status", "downloads", "captcha_done", "reset", "menu"): out.write(f"  <- {kind}: { {k: v for k, v in ev.items() if k != 'type'} }\n")

def wait_for(pred, timeout=8):
    deadline = time.time() + timeout
    while time.time() < deadline:
        try:
            ev = events.get(timeout=0.1)
        except queue.Empty:
            continue
        if pred(ev): return ev
    raise AssertionError("timed out waiting for event")

def say(msg): sys.__stdout__.write(msg + "\n")
def idle(): wait_for(lambda e: e["type"] == "busy" and not e["busy"])

ui = RecordingUI()
worker = gui.Worker(ui)
worker.start()
sys.stdout = gui.StdoutTee(sys.stdout, ui)
FP = FAKE_PW.browser.page

say("\n== 1. login (wrong captcha, then right) -> menu ==")
worker.post("login", username="me", password="pw", headed=False, solver="manual", preview=True)
ev = wait_for(lambda e: e["type"] == "captcha"); assert ev["allow_skip"] is False
worker.captcha_answers.put("WRONG")
wait_for(lambda e: e["type"] == "captcha")
worker.captcha_answers.put(CAPTCHA_ANSWER)
wait_for(lambda e: e["type"] == "status" and e["connected"])
wait_for(lambda e: e["type"] == "menu")
assert FP.login_attempts == 2
idle()

say("\n== 2. open Buyer mode -> page (generic form) ==")
worker.post("open_mode", mode="buyer")
pg = wait_for(lambda e: e["type"] == "page")
assert pg["mode"] == "buyer" and FP.url == BUYER_URL and len(pg["controls"]) == 7
idle()

say("\n== 2b. Khasra mode searches a deduplicated list across selected years ==")
worker.post("open_mode", mode="khasra")
pg = wait_for(lambda e: e["type"] == "page")
assert pg["mode"] == "khasra" and FP.role == "khasra" and pg["controls"][3]["label"] == "Khasra No."
idle()
FP.khasra_searches.clear()
worker.post("run_khasra_search", params={"district": "BAGESHWAR", "sro": "BAGESHWAR", "from_year": "2022",
                                         "to_year": "2023", "khasra_numbers": " 12, 13\n12 "})
while True:
    ev = wait_for(lambda e: e["type"] in ("captcha", "khasra_done"), timeout=15)
    if ev["type"] == "captcha": worker.captcha_answers.put(CAPTCHA_ANSWER)
    else: break
assert ev["partial"] is False and ev["header"] == ["Khasra No.", "Year", "Sr", "Buyer", "Doc"]
assert FP.khasra_searches == [("2022", "12"), ("2022", "13"), ("2023", "12"), ("2023", "13")], FP.khasra_searches
assert len(ev["rows"]) == 12 and ev["params"]["khasra_numbers"] == ["12", "13"]
idle()

say("\n== 3. open Seller mode -> Seller Wise page ==")
worker.post("open_mode", mode="seller")
pg = wait_for(lambda e: e["type"] == "page")
assert pg["mode"] == "seller" and FP.url == SELLER_URL and FP.role == "seller"
idle()

say("\n== 4. dropdown change on the live form still works ==")
worker.post("set_value", selector="#MainContent_ddlDistrict", value={"value": "2", "label": "BAGESHWAR", "field": "District"}, kind="select")
wait_for(lambda e: e["type"] == "page"); idle()

say("\n== 5. generic Search click with captcha (rejected once) ==")
worker.post("click", selector="#MainContent_btnSearch", values={"#MainContent_txtBuyer": "आशुतोष जौहरी"})
wait_for(lambda e: e["type"] == "captcha"); worker.captcha_answers.put("nope")
wait_for(lambda e: e["type"] == "captcha"); worker.captcha_answers.put(CAPTCHA_ANSWER)
pg = wait_for(lambda e: e["type"] == "page"); assert pg["tables"]; idle()

say("\n== 6. back to menu, open report mode ==")
worker.post("menu"); wait_for(lambda e: e["type"] == "menu"); idle()
worker.post("open_mode", mode="report")
pg = wait_for(lambda e: e["type"] == "page"); assert pg["mode"] == "report" and FP.url == BUYER_URL; idle()

say("\n== 7. run report: buyer + seller, 2022-2023, two result pages each, captcha per search ==")
FP.searches.clear()
worker.post("run_report", params={"roles": ["buyer", "seller"], "district": "BAGESHWAR", "sro": "BAGESHWAR",
                                  "from_year": "2023", "to_year": "2022", "name": "आशुतोष"})
done = None
for _ in range(40):
    ev = wait_for(lambda e: e["type"] in ("captcha", "report_done"), timeout=15)
    if ev["type"] == "captcha":
        assert "Search" in ev["prompt"]; worker.captcha_answers.put(CAPTCHA_ANSWER)
    else:
        done = ev; break
assert done and not done["partial"], done
assert [s[:2] for s in FP.searches] == [("buyer", "2022"), ("buyer", "2023"), ("seller", "2022"), ("seller", "2023")], FP.searches
assert all(s[2] == "आशुतोष" for s in FP.searches)
for role in ("buyer", "seller"):
    sheet = done["sheets"][role]
    assert sheet["header"] == ["ID", "Year", "Sr", "Buyer", "Doc", "Relation", "Relative"], sheet["header"]
    assert len(sheet["rows"]) == 6, sheet["rows"]                       # 2 years x (2 + 1 rows), pager rows dropped
    R = role[0].upper()
    assert sheet["rows"][0] == [f"{R}2022-0001", "2022", "1", f"{role}-2022-p1-A", "Download Document", "", ""]
    assert sheet["rows"][2] == [f"{R}2022-0003", "2022", "3", f"{role}-2022-p2-C", "Download Document", "", ""]
    assert sheet["rows"][3][:2] == [f"{R}2023-0001", "2023"]
    assert len({r[0] for r in sheet["rows"]}) == 6                     # IDs unique within the sheet
assert done["sheets"]["buyer"]["party_column"] == "Buyer" and done["sheets"]["buyer"]["party_index"] == 3
assert done["sheets"]["seller"]["party_index"] == -1                    # mock table has no seller column
# district/SRO were re-applied by label on each role page
assert FP.filled["#MainContent_ddlDistrict"] == {"label": "BAGESHWAR"} and FP.filled["#MainContent_ddlSRO"] == {"label": "BAGESHWAR"}
idle()

say("\n== 8. export: xlsx if openpyxl is available, else csv; filtered selection ==")
tmp = Path(tempfile.mkdtemp())
info = worker.export_report(str(tmp / "r.xlsx"), {"buyer": [0, 2], "seller": None})
say(f"  export -> {info}")
assert info["counts"] == {"Buyer": 2, "Seller": 6}
assert all(Path(f).exists() for f in info["files"])
if info["format"] == "xlsx":
    import openpyxl
    wb = openpyxl.load_workbook(info["path"])
    assert wb.sheetnames == ["Buyer", "Seller", "Summary"], wb.sheetnames
    ws = wb["Buyer"]; rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("ID", "Year", "Sr", "Buyer", "Doc", "Relation", "Relative") and len(rows) == 3, rows
    assert rows[1][0] == "B2022-0001" and rows[1][3] == "buyer-2022-p1-A", rows
else:
    text = Path(info["files"][0]).read_text(encoding="utf-8-sig")
    assert "buyer-2022-p1-A" in text and "buyer-2022-p1-B" not in text
shutil.rmtree(tmp)

say("\n== 9. cancel mid-report -> partial results ==")
worker.post("run_report", params={"roles": ["buyer"], "district": "BAGESHWAR", "sro": "BAGESHWAR",
                                  "from_year": "2021", "to_year": "2024", "name": "x"})
wait_for(lambda e: e["type"] == "captcha"); worker.captcha_answers.put(CAPTCHA_ANSWER)          # year 2021 completes
wait_for(lambda e: e["type"] == "report_progress" and str(e["message"]).startswith("done"))
gui.Api(worker).cancel_report()                                                                # Stop pressed
ev = wait_for(lambda e: e["type"] in ("report_done",), timeout=15)
assert ev["partial"] is True and 3 <= len(ev["sheets"]["buyer"]["rows"]) <= 6, ev["sheets"]["buyer"]["rows"]
idle()

say("\n== 9b. settings.json remembers the report fields; report mode re-applies District/SRO/year ==")
saved = json.loads(gui.SETTINGS_FILE.read_text(encoding="utf-8"))
assert saved["report"] == {"district": "BAGESHWAR", "sro": "BAGESHWAR", "from_year": 2021, "to_year": 2024, "name": "x", "roles": ["buyer"]}, saved
api0 = gui.Api(worker)
with patch.dict(os.environ, {"CAPTCHA_SOLVER": "", "UK_EREG_USERNAME": "", "UK_EREG_PASSWORD": ""}):
    api0.save_prefs({"solver": "openai", "download_dir": "/tmp/docs", "bogus": 1})
    d = api0.defaults()
    assert d["default_solver"] == "openai" and d["download_dir"] == "/tmp/docs" and "bogus" not in d["settings"]
    api0.save_prefs({"remember_credentials": True, "username": "saved-user", "password": "saved-pass"})
    d = api0.defaults()
    assert d["username"] == "saved-user" and d["password"] == "saved-pass" and d["remember_credentials"]
    assert "credentials" not in d["settings"]
    os.environ["UK_EREG_USERNAME"] = "env-user"
    os.environ["UK_EREG_PASSWORD"] = "env-pass"
    d = api0.defaults()
    assert d["username"] == "env-user" and d["password"] == "env-pass"
    os.environ["UK_EREG_USERNAME"] = ""
    os.environ["UK_EREG_PASSWORD"] = ""
    api0.save_prefs({"remember_credentials": False})
    d = api0.defaults()
    saved_credentials = json.loads(gui.SETTINGS_FILE.read_text(encoding="utf-8"))["credentials"]
    assert not d["remember_credentials"] and saved_credentials == {"remember": False, "username": "", "password": ""}
FP.filled.clear()
worker.post("open_mode", mode="report")
pg = wait_for(lambda e: e["type"] == "page"); idle()
assert FP.filled["#MainContent_ddlDistrict"] == {"label": "BAGESHWAR"} and FP.filled["#MainContent_ddlYear"] == {"label": "2021"}, FP.filled

say("\n== 9c. download documents for selected report rows (re-search, walk to page 2, click each row's link) ==")
worker.post("run_report", params={"roles": ["buyer"], "district": "BAGESHWAR", "sro": "BAGESHWAR",
                                  "from_year": "2022", "to_year": "2023", "name": "आशुतोष"})
while True:
    ev = wait_for(lambda e: e["type"] in ("captcha", "report_done"), timeout=15)
    if ev["type"] == "captcha": worker.captcha_answers.put(CAPTCHA_ANSWER)
    else: break
idle()
assert "meta" not in ev and len(worker.report["meta"]["buyer"]) == 6
tmp = Path(tempfile.mkdtemp())
FP.doc_clicks.clear()
# rows 0 (2022 p1), 2 (2022 p2), 5 (2023 p2) of the buyer sheet
worker.post("download_rows", selection={"buyer": [0, 2, 5]}, out_dir=str(tmp))
while True:
    ev = wait_for(lambda e: e["type"] in ("captcha", "downloads"), timeout=15)
    if ev["type"] == "captcha": worker.captcha_answers.put(CAPTCHA_ANSWER)
    elif ev["done"]: break
idle()
assert len(ev["files"]) == 3, ev
names = sorted(Path(f).name for f in ev["files"])
assert names == ["B2022-0001.pdf", "B2022-0003.pdf", "B2023-0003.pdf"], names            # <row ID>.pdf
assert [(c[1], c[2], c[3]) for c in FP.doc_clicks] == [("2022", 1, "buyer-2022-p1-A"), ("2022", 2, "buyer-2022-p2-C"), ("2023", 2, "buyer-2023-p2-C")], FP.doc_clicks
assert all(Path(f).read_bytes().startswith(b"%PDF") for f in ev["files"])
shutil.rmtree(tmp)

say("\n== 10. automatic captcha solvers: a good one logs in silently, a bad one falls back to manual ==")
import captcha_solvers  # noqa: E402
captcha_solvers.SOLVERS["fake_good"] = {"label": "Fake good", "fn": lambda png: CAPTCHA_ANSWER, "check": lambda: (True, "")}
captcha_solvers.SOLVERS["fake_bad"] = {"label": "Fake bad", "fn": lambda png: "WRONG", "check": lambda: (True, "")}
worker.post("logout"); wait_for(lambda e: e["type"] == "reset"); idle()
FP.__init__()                                                  # fresh fake page (login state)
FAKE_PW.browser.page = FP
worker.post("login", username="me", password="pw", headed=False, solver="fake_good", preview=False)
ev = wait_for(lambda e: e["type"] in ("captcha", "menu"))
assert ev["type"] == "menu", "a working solver must not prompt"
assert FP.login_attempts == 1
idle()
worker.post("logout"); wait_for(lambda e: e["type"] == "reset"); idle()
FP.__init__(); FAKE_PW.browser.page = FP
worker.post("login", username="me", password="pw", headed=False, solver="fake_bad", preview=False)
ev = wait_for(lambda e: e["type"] == "captcha", timeout=10)     # after two rejected solver answers -> manual box
assert FP.login_attempts == 2, FP.login_attempts
worker.captcha_answers.put(CAPTCHA_ANSWER)
wait_for(lambda e: e["type"] == "menu"); idle()
api = gui.Api(worker)
assert any(s["id"] == "manual" and s["available"] for s in api.defaults()["solvers"])

say("\n== 10b. Chromium missing on a new machine -> the app runs `playwright install chromium` itself ==")
import io  # noqa: E402
class FakeProc:
    def __init__(self, cmd, **k):
        FakeProc.cmd = cmd
        out = ("Downloading Chromium 131.0.1 (playwright build v1148) from https://example/chromium.zip\r"
               "|         |   0% of 165.6 MiB\r|====     |  50% of 165.6 MiB\r|=========| 100% of 165.6 MiB\n"
               "Chromium 131.0.1 (playwright build v1148) downloaded to /x/ms-playwright/chromium-1148\n")
        self.stdout = io.StringIO(out)
    def wait(self): return 0
real_popen = gui.subprocess.Popen
gui.subprocess.Popen = FakeProc
worker.post("logout"); wait_for(lambda e: e["type"] == "reset"); idle()
FP.__init__(); FAKE_PW.browser.page = FP; FAKE_PW.fail_launch_once = True
launches_before = FAKE_PW.launches
worker.post("login", username="me", password="pw", headed=False, solver="fake_good", preview=False)
ev = wait_for(lambda e: e["type"] == "log" and "downloading it now" in e["message"])
ev = wait_for(lambda e: e["type"] == "busy" and str(e.get("what", "")).endswith("100%"))
wait_for(lambda e: e["type"] == "log" and e["message"] == "Chromium downloaded.")
wait_for(lambda e: e["type"] == "menu"); idle()
assert FAKE_PW.launches == launches_before + 2, FAKE_PW.launches            # failed once, then succeeded
assert FakeProc.cmd[-2:] == ["install", "chromium"], FakeProc.cmd
gui.subprocess.Popen = real_popen

say("\n== 10c. relation parsing (Relation / Relative columns) ==")
pr = gui.parse_relation
assert pr("अनुराग गुप्ता S/O ठाकुर दास गुप्ता , ,") == ("S/O", "ठाकुर दास गुप्ता")
assert pr("मनप्रीत कौर W/O दलजीत सिंह , ,") == ("W/O", "दलजीत सिंह")
assert pr("अनुराग गुप्ता एच0यू0एफ0 द्वारा कर्ता श्री अनुराग गुप्ता S/O रमेश चन्द गुप्ता , ,") == ("S/O", "रमेश चन्द गुप्ता")
assert pr("राम प्रसाद पुत्र श्याम लाल निवासी देहरादून") == ("पुत्र", "श्याम लाल")
assert pr("सीता देवी पत्नी मोहन सिंह, देहरादून") == ("पत्नी", "मोहन सिंह")
assert pr("रीता पुत्री हरि") == ("पुत्री", "हरि")
assert pr("kumar s/0 ramesh") == ("S/O", "ramesh") and pr("HUF") == ("", "") and pr("") == ("", "")
assert gui.party_column_index(["Details", "Seller/First Party", "Buyer/Second Party", "Buyer/Second Gender", "SRO Name"], "buyer") == 2
assert gui.party_column_index(["Details", "Seller/First Party", "Buyer/Second Party", "Buyer/Second Gender", "SRO Name"], "seller") == 1
assert gui.party_column_index(["Sr", "Doc"], "buyer") is None

say("\n== 11. Api validation + shutdown ==")
assert api.run_report({"from_year": "abc"})["ok"] is False
assert api.login("", "x")["ok"] is False
worker.shutdown()
assert not worker.is_alive() and FAKE_PW.browser.closed and FAKE_PW.stopped
sys.stdout = sys.__stdout__
say("\nALL FLOW TESTS PASSED")
